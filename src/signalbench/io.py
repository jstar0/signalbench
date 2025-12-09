from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import tempfile
import unicodedata
import zipfile
from collections.abc import Iterable, Iterator, Mapping, Sequence
from pathlib import Path
from typing import Any

from .config import DEFAULT_CONFIG
from .models import PREPARED_COLUMNS, QUALITY_LABELS, PreparedRecord, SourceRecord
from .quality import classify_quality, text_fingerprint

SUPPORTED_SUFFIXES = {".csv", ".tsv", ".xlsx"}
_EXCLUDED_DIR_NAMES = {".git", ".venv", "__pycache__", "outputs", ".pytest_cache"}
_HEADER_PUNCTUATION_RE = re.compile(r"[\s_\-./:：()（）\[\]{}]+")

# SignalBench intentionally materialises prepared records for deterministic
# duplicate grouping, resume, and CSV/XLSX artifacts. These limits make that
# v1 memory contract explicit before a large upload is hashed or expanded.
MAX_INPUT_FILES = 100
MAX_INPUT_BYTES = 16 * 1024 * 1024
MAX_COLUMNS = 128
MAX_RECORDS = 10_000
MAX_TEXT_CHARS = 8_192
MAX_METADATA_CHARS = 512


class InputContractError(ValueError):
    """Raised when an input table cannot be mapped to the signal schema."""


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def stable_digest(value: str, length: int = 20) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()[:length]


def normalize_header(value: object) -> str:
    text = "" if value is None else unicodedata.normalize("NFKC", str(value))
    text = text.strip().casefold()
    return _HEADER_PUNCTUATION_RE.sub("", text)


def discover_inputs(inputs: Sequence[Path]) -> list[Path]:
    """Expand files/directories into a deterministic list of supported inputs."""

    candidates: list[Path] = []
    for item in inputs:
        path = item.expanduser().resolve()
        if not path.exists():
            raise FileNotFoundError(f"Input does not exist: {item}")
        if path.is_dir():
            candidates.extend(
                child
                for child in path.rglob("*")
                if child.is_file()
                and child.suffix.lower() in SUPPORTED_SUFFIXES
                and not any(part in _EXCLUDED_DIR_NAMES for part in child.relative_to(path).parts)
            )
        elif path.suffix.lower() in SUPPORTED_SUFFIXES:
            candidates.append(path)
        else:
            raise InputContractError(
                f"Unsupported input format {path.suffix!r}; use CSV, TSV, or XLSX."
            )
    unique = {path.resolve() for path in candidates}
    if not unique:
        raise InputContractError("No CSV, TSV, or XLSX inputs were found.")
    return sorted(unique, key=lambda path: str(path).casefold())


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value)


def _find_column(headers: Sequence[str], aliases: Sequence[str], *, required: bool) -> str | None:
    normalized = {normalize_header(header): header for header in headers}
    for alias in aliases:
        match = normalized.get(normalize_header(alias))
        if match is not None:
            return match
    if required:
        raise InputContractError(
            "Could not find a text column. Supported names: " + ", ".join(aliases)
        )
    return None


def _fallback_platform(source_sha256: str) -> str:
    """Return a stable report-safe label when no explicit platform exists."""

    return f"source-{source_sha256[:12]}"


def _record_from_row(
    *,
    path: Path,
    source_sha256: str,
    sheet: str,
    row_number: int,
    row: Mapping[str, object],
    columns: Mapping[str, str | None],
) -> SourceRecord | None:
    text_col = columns["text"]
    assert text_col is not None
    text_raw = _cell_text(row.get(text_col, ""))
    # Ignore genuinely empty spreadsheet rows; retain a row with a visible value
    # so the quality gate can explain why it is not analyzable.
    if not text_raw.strip() and not any(_cell_text(value).strip() for value in row.values()):
        return None
    if len(text_raw) > MAX_TEXT_CHARS:
        raise InputContractError(
            f"{path.name}:{row_number} text exceeds the {MAX_TEXT_CHARS}-character v1 limit."
        )
    platform_col = columns["platform"]
    author_col = columns["author"]
    id_col = columns["record_id"]
    platform = _cell_text(row.get(platform_col, "")) if platform_col else ""
    platform = platform.strip() or _fallback_platform(source_sha256)
    author = _cell_text(row.get(author_col, "")) if author_col else ""
    author_hash = stable_digest(author, 12) if author.strip() else ""
    explicit_id = _cell_text(row.get(id_col, "")) if id_col else ""
    for field, value in (("platform", platform), ("author", author), ("record id", explicit_id)):
        if len(value) > MAX_METADATA_CHARS:
            raise InputContractError(
                f"{path.name}:{row_number} {field} exceeds the {MAX_METADATA_CHARS}-character v1 limit."
            )
    seed = "\x1f".join(
        [source_sha256, sheet, str(row_number), explicit_id.strip(), text_raw]
    )
    record_id = f"sb_{stable_digest(seed)}"
    return SourceRecord(
        record_id=record_id,
        source_file=path.name,
        source_sheet=sheet,
        source_row=row_number,
        platform=platform,
        author_hash=author_hash,
        text_raw=text_raw,
    )


def _resolve_columns(headers: Sequence[str], config: Mapping[str, Any]) -> dict[str, str | None]:
    if len(headers) > MAX_COLUMNS:
        raise InputContractError(f"Input has {len(headers)} columns; v1 supports at most {MAX_COLUMNS}.")
    if any(len(header) > MAX_METADATA_CHARS for header in headers):
        raise InputContractError(f"Input has a header longer than the {MAX_METADATA_CHARS}-character v1 limit.")
    schema = config["schema"]
    return {
        "text": _find_column(headers, schema["text"], required=True),
        "platform": _find_column(headers, schema["platform"], required=False),
        "author": _find_column(headers, schema["author"], required=False),
        "record_id": _find_column(headers, schema["record_id"], required=False),
    }


def _iter_csv_rows(path: Path) -> tuple[list[str], Iterator[tuple[int, dict[str, str]]]]:
    # Read a small sample only to infer the delimiter; the actual file is opened
    # again by the iterator, keeping memory bounded for large exports.
    encoding: str | None = None
    for candidate in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            with path.open("r", encoding=candidate, newline="") as handle:
                handle.read(4096)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    if encoding is None:
        raise InputContractError(f"Unable to decode {path.name} as UTF-8 or GB18030.")

    with path.open("r", encoding=encoding, newline="") as handle:
        sample = handle.read(4096)
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel

    def iterator() -> Iterator[tuple[int, dict[str, str]]]:
        with path.open("r", encoding=encoding, newline="") as handle:
            reader = csv.DictReader(handle, dialect=dialect)
            if reader.fieldnames is None:
                raise InputContractError(f"{path.name} has no header row.")
            for row_number, row in enumerate(reader, start=2):
                yield row_number, {str(key): _cell_text(value) for key, value in row.items() if key is not None}

    with path.open("r", encoding=encoding, newline="") as handle:
        reader = csv.reader(handle, dialect=dialect)
        try:
            headers = [str(value) for value in next(reader)]
        except StopIteration as exc:
            raise InputContractError(f"{path.name} is empty.") from exc
    return headers, iterator()


def _iter_xlsx_sheets(path: Path) -> Iterator[tuple[str, list[str], Iterator[tuple[int, dict[str, str]]]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency is declared in pyproject
        raise RuntimeError("XLSX support requires openpyxl; run `uv sync` first.") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            try:
                raw_headers = next(rows)
            except StopIteration:
                continue
            headers = [_cell_text(value) for value in raw_headers]
            if not any(header.strip() for header in headers):
                continue

            def iterator(
                remaining_rows: Iterator[tuple[object, ...]],
                header_values: list[str],
            ) -> Iterator[tuple[int, dict[str, str]]]:
                for row_number, values in enumerate(remaining_rows, start=2):
                    padded = list(values) + [None] * max(0, len(header_values) - len(values))
                    yield row_number, {
                        header_values[index]: _cell_text(padded[index])
                        for index in range(len(header_values))
                        if header_values[index]
                    }

            yield worksheet.title, headers, iterator(rows, headers)
    finally:
        workbook.close()


def load_source_records(
    paths: Sequence[Path], config: Mapping[str, Any] | None = None
) -> tuple[list[SourceRecord], dict[str, Any]]:
    """Load supported tables and return rows plus an auditable input manifest."""

    resolved_config = config or DEFAULT_CONFIG
    records: list[SourceRecord] = []
    manifest_files: list[dict[str, Any]] = []
    seen_ids: set[str] = set()
    inputs_found = discover_inputs(paths)
    if len(inputs_found) > MAX_INPUT_FILES:
        raise InputContractError(f"Input expands to {len(inputs_found)} files; v1 supports at most {MAX_INPUT_FILES}.")

    def append_record(record: SourceRecord | None, *, path: Path, row_number: int) -> bool:
        if record is None:
            return False
        if len(records) >= MAX_RECORDS:
            raise InputContractError(f"Input exceeds the {MAX_RECORDS}-record v1 limit at {path.name}:{row_number}.")
        if record.record_id in seen_ids:
            raise InputContractError(f"Stable record id collision at {path.name}:{row_number}.")
        seen_ids.add(record.record_id)
        records.append(record)
        return True

    for path in inputs_found:
        file_bytes = path.stat().st_size
        if file_bytes > MAX_INPUT_BYTES:
            raise InputContractError(
                f"{path.name} is {file_bytes} bytes; v1 supports files up to {MAX_INPUT_BYTES} bytes."
            )
        source_sha256 = sha256_file(path)
        file_entry: dict[str, Any] = {
            "file": path.name,
            "format": path.suffix.lower().lstrip("."),
            "sha256": source_sha256,
            "sheets": [],
        }
        if path.suffix.lower() in {".csv", ".tsv"}:
            headers, rows = _iter_csv_rows(path)
            sheet = "__csv__"
            columns = _resolve_columns(headers, resolved_config)
            sheet_entry = {
                "name": sheet,
                "columns": headers,
                "detected_columns": columns,
                "rows": 0,
            }
            for row_number, row in rows:
                record = _record_from_row(
                    path=path,
                    source_sha256=source_sha256,
                    sheet=sheet,
                    row_number=row_number,
                    row=row,
                    columns=columns,
                )
                if append_record(record, path=path, row_number=row_number):
                    sheet_entry["rows"] += 1
            file_entry["sheets"].append(sheet_entry)
        else:
            for sheet, headers, rows in _iter_xlsx_sheets(path):
                columns = _resolve_columns(headers, resolved_config)
                sheet_entry = {
                    "name": sheet,
                    "columns": headers,
                    "detected_columns": columns,
                    "rows": 0,
                }
                for row_number, row in rows:
                    record = _record_from_row(
                        path=path,
                        source_sha256=source_sha256,
                        sheet=sheet,
                        row_number=row_number,
                        row=row,
                        columns=columns,
                    )
                    if append_record(record, path=path, row_number=row_number):
                        sheet_entry["rows"] += 1
                file_entry["sheets"].append(sheet_entry)
        manifest_files.append(file_entry)
    manifest = {
        "schema_version": 1,
        "record_count": len(records),
        "files": manifest_files,
    }
    if not records:
        raise InputContractError("Inputs contained no non-empty data rows.")
    return records, manifest


def ensure_parent(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)


def atomic_write_text(path: Path, content: str) -> None:
    """Write a file atomically so an interrupted run cannot corrupt a checkpoint."""

    ensure_parent(path)
    fd, temp_name = tempfile.mkstemp(prefix=f".{path.name}.", dir=str(path.parent))
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="") as handle:
            handle.write(content)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temp_name, path)
    except Exception:
        try:
            os.unlink(temp_name)
        except FileNotFoundError:
            pass
        raise


def write_json(path: Path, payload: Mapping[str, Any]) -> None:
    atomic_write_text(path, json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n")


def write_prepared_csv(path: Path, records: Iterable[PreparedRecord]) -> None:
    rows = [record.to_row() for record in records]
    ensure_parent(path)
    fd, temp_name = tempfile.mkstemp(prefix=f".{path.name}.", dir=str(path.parent))
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=PREPARED_COLUMNS, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temp_name, path)
    except Exception:
        try:
            os.unlink(temp_name)
        except FileNotFoundError:
            pass
        raise


def read_prepared_csv(path: Path, config: Mapping[str, Any] | None = None) -> list[PreparedRecord]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames is None:
            raise InputContractError(f"{path.name} has no header row.")
        missing = {"record_id", "text_normalized", "quality_label"} - set(reader.fieldnames)
        if missing:
            raise InputContractError(f"Prepared file is missing columns: {', '.join(sorted(missing))}")
        records: list[PreparedRecord] = []
        for row_number, row in enumerate(reader, start=2):
            if len(records) >= MAX_RECORDS:
                raise InputContractError(
                    f"Prepared file exceeds the {MAX_RECORDS}-record v1 limit at {path.name}:{row_number}."
                )
            records.append(PreparedRecord.from_row(row))
    for record in records:
        if not record.record_id:
            raise InputContractError("Prepared file contains an empty record_id.")
        if record.quality_label not in QUALITY_LABELS:
            raise InputContractError(
                f"Prepared file has unsupported quality_label for {record.record_id}: {record.quality_label}"
            )
        if record.classification_status not in {"pending", "classified", "not_applicable"}:
            raise InputContractError(
                f"Prepared file has unsupported classification_status for {record.record_id}: {record.classification_status}"
            )
        expected_fingerprint = text_fingerprint(record.text_normalized) if record.text_normalized else ""
        if record.text_fingerprint != expected_fingerprint:
            raise InputContractError(f"Text fingerprint mismatch for {record.record_id}.")
        # Recompute the gate rather than trusting the column. Verifying only the
        # fingerprint let an edited row flip `invalid` to `valid` and reach the
        # classifier, which is the one thing the quality gate exists to prevent.
        expected_label = classify_quality(record.text_normalized, config).label
        if record.quality_label != expected_label:
            raise InputContractError(
                f"Quality label for {record.record_id} does not match its text: "
                f"file says {record.quality_label}, rules say {expected_label}."
            )
        expected_status = "pending" if expected_label == "valid" else "not_applicable"
        if record.classification_status == "not_applicable" and expected_label == "valid":
            raise InputContractError(
                f"Record {record.record_id} is valid but marked {record.classification_status}."
            )
        if record.classification_status != "classified" and expected_label != "valid" and record.classification_status != expected_status:
            raise InputContractError(
                f"Record {record.record_id} is {expected_label} but marked {record.classification_status}."
            )
        if record.duplicate_count < 1 or not (1 <= record.duplicate_rank <= record.duplicate_count):
            raise InputContractError(f"Invalid duplicate rank/count for {record.record_id}.")
    if len({record.record_id for record in records}) != len(records):
        raise InputContractError("Prepared file contains duplicate record_id values.")
    return records


def write_rows_xlsx(path: Path, sheets: Mapping[str, Sequence[Mapping[str, Any]]]) -> None:
    """Write machine-readable rows to an XLSX workbook using write-only mode."""

    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover - dependency is declared in pyproject
        raise RuntimeError("XLSX export requires openpyxl; run `uv sync` first.") from exc

    ensure_parent(path)
    fd, temp_name = tempfile.mkstemp(prefix=f".{path.stem}.", suffix=".xlsx", dir=str(path.parent))
    os.close(fd)
    try:
        workbook = Workbook(write_only=True)
        # openpyxl otherwise places the current time in the workbook metadata.
        # A generated spreadsheet is part of this project's reproducible
        # evidence, so it must not change just because it was written later.
        from datetime import UTC, datetime

        fixed_timestamp = datetime(1980, 1, 1, tzinfo=UTC)
        workbook.properties.created = fixed_timestamp
        workbook.properties.modified = fixed_timestamp
        workbook.properties.lastModifiedBy = "SignalBench"
        used_titles: set[str] = set()
        for sheet_name, rows in sheets.items():
            base_title = re.sub(r"[:\\/?*\[\]]", "_", str(sheet_name))[:31] or "Sheet1"
            title = base_title
            suffix = 2
            while title in used_titles:
                suffix_text = f"_{suffix}"
                title = f"{base_title[: 31 - len(suffix_text)]}{suffix_text}"
                suffix += 1
            used_titles.add(title)
            worksheet = workbook.create_sheet(title=title)
            row_list = list(rows)
            columns: list[str] = []
            seen: set[str] = set()
            for row in row_list:
                for key in row:
                    if key not in seen:
                        columns.append(key)
                        seen.add(key)
            if not columns:
                columns = ["value"]
            worksheet.append(columns)
            for row in row_list:
                values: list[object] = []
                for column in columns:
                    value = row.get(column, "")
                    if isinstance(value, (dict, list)):
                        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
                    values.append(value)
                worksheet.append(values)
        workbook.save(temp_name)
        workbook.close()
        _canonicalize_xlsx_archive(Path(temp_name))
        os.replace(temp_name, path)
    except Exception:
        try:
            os.unlink(temp_name)
        except FileNotFoundError:
            pass
        raise


def _canonicalize_xlsx_archive(path: Path) -> None:
    """Rewrite ZIP headers in a stable order and timestamp for hashable XLSX."""

    with zipfile.ZipFile(path, "r") as source:
        members = []
        for name in sorted(source.namelist()):
            payload = source.read(name)
            if name == "docProps/core.xml":
                # openpyxl rewrites `modified` while saving, after properties
                # are assigned. Normalize that final metadata byte sequence too.
                payload = re.sub(
                    rb'<dcterms:modified[^>]*>.*?</dcterms:modified>',
                    b'<dcterms:modified xsi:type="dcterms:W3CDTF">1980-01-01T00:00:00Z</dcterms:modified>',
                    payload,
                )
            members.append((name, payload))
    canonical = path.with_name(f".{path.name}.canonical")
    try:
        with zipfile.ZipFile(canonical, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target:
            for name, payload in members:
                info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                target.writestr(info, payload, compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)
        os.replace(canonical, path)
    except Exception:
        try:
            canonical.unlink()
        except FileNotFoundError:
            pass
        raise


def write_prepared_xlsx(path: Path, records: Iterable[PreparedRecord]) -> None:
    rows = [record.to_row() for record in records]
    write_rows_xlsx(path, {"prepared": rows})
